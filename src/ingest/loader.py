from pathlib import Path
from typing import Any

from docx import Document
from pypdf import PdfReader

from ingest.legacy import read_with_libreoffice
from ingest.office import read_pptx, read_xls, read_xlsx


def _read_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


def _read_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    parts: list[str] = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            parts.append(text)
    return "\n\n".join(parts)


def _read_docx(path: Path) -> str:
    doc = Document(str(path))
    parts: list[str] = []

    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)

    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append("\t".join(cells))

    return "\n\n".join(parts)


def _read_doc(path: Path) -> str:
    return read_with_libreoffice(path)


def _read_ppt(path: Path) -> str:
    return read_with_libreoffice(path)


READERS = {
    ".txt": _read_txt,
    ".md": _read_txt,
    ".pdf": _read_pdf,
    ".docx": _read_docx,
    ".doc": _read_doc,
    ".pptx": read_pptx,
    ".ppt": _read_ppt,
    ".xlsx": read_xlsx,
    ".xls": read_xls,
}


def ingest_file(path: Path) -> str:
    reader = READERS.get(path.suffix.lower())
    if reader is None:
        raise ValueError(f"Unsupported file type: {path.suffix}")
    text = reader(path).strip()
    if not text:
        raise ValueError(f"No text extracted from {path.name}")
    return text


def ingest_all(config: dict[str, Any], root: Path) -> int:
    raw_dir = root / config["ingest"]["raw_dir"]
    out_dir = root / config["ingest"]["processed_dir"]
    exts = set(config["ingest"]["supported_extensions"])

    raw_dir.mkdir(parents=True, exist_ok=True)
    out_dir.mkdir(parents=True, exist_ok=True)

    count = 0
    for path in sorted(raw_dir.rglob("*")):
        if not path.is_file() or path.suffix.lower() not in exts:
            continue
        try:
            text = ingest_file(path)
        except Exception as exc:
            print(f"  SKIP {path.name}: {exc}")
            continue

        rel = path.relative_to(raw_dir)
        out_path = out_dir / rel.with_suffix(".txt")
        out_path.parent.mkdir(parents=True, exist_ok=True)
        out_path.write_text(text, encoding="utf-8")
        print(f"  OK   {path.name} -> {out_path.relative_to(root)}")
        count += 1

    return count
